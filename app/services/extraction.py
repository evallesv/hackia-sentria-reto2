"""Servicio de extracción estructurada de documentos (PDF/XLSX) con trazabilidad (T02)."""

import re
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from fastapi import HTTPException
from openpyxl import load_workbook
from pypdf import PdfReader

from app.models import (
    DocumentKind,
    Evidence,
    ExtractedItem,
    ExtractedTariff,
    ExtractionSnapshot,
    StoredDocument,
)

MAX_PDF_PAGES = 30
MAX_XLSX_ROWS = 1000
MIN_TEXT_DENSITY_CHARS = 20


def extract_pdf_pages(file_path: Path, max_pages: int = MAX_PDF_PAGES) -> list[tuple[int, str]]:
    """Extrae texto página por página usando pypdf; previene alucinaciones en escaneados."""
    try:
        reader = PdfReader(str(file_path))
    except Exception as e:
        raise HTTPException(400, f"Error al leer el archivo PDF: {e}") from None

    if reader.is_encrypted:
        raise HTTPException(400, "El documento PDF está protegido o cifrado y no puede procesarse")

    total_pages = len(reader.pages)
    if total_pages > max_pages:
        raise HTTPException(
            400, f"El documento PDF excede el límite máximo de {max_pages} páginas ({total_pages})"
        )

    pages_text: list[tuple[int, str]] = []
    total_chars = 0

    for idx, page in enumerate(reader.pages):
        page_num = idx + 1
        raw_text = page.extract_text() or ""
        clean_lines = [re.sub(r"[ \t]+", " ", line).strip() for line in raw_text.splitlines()]
        clean_text = "\n".join(line for line in clean_lines if line)
        total_chars += len(clean_text)
        pages_text.append((page_num, clean_text))

    if total_chars < MIN_TEXT_DENSITY_CHARS:
        # Documento sin texto legible / imagen escaneada: jamás inventar texto
        return [(1, "[DOCUMENTO ESCANEADO O SIN TEXTO DIGITAL LEGIBLE - REQUIERE OCR]")]

    return pages_text


def extract_xlsx_tariffs(file_path: Path, max_rows: int = MAX_XLSX_ROWS) -> list[dict]:
    """Extrae tarifas de un tarifario en XLSX en modo sólo lectura y sin evaluar fórmulas."""
    try:
        wb = load_workbook(filename=str(file_path), read_only=True, data_only=False)
    except Exception as e:
        raise HTTPException(400, f"Error al abrir la hoja de cálculo XLSX: {e}") from None

    sheet = wb.active
    if not sheet:
        raise HTTPException(400, "La hoja de cálculo XLSX no contiene hojas activas")

    rows = list(sheet.iter_rows(values_only=True, max_row=max_rows))
    wb.close()

    if not rows:
        raise HTTPException(400, "La hoja de cálculo está vacía")

    header_row = [str(cell).strip().lower() if cell is not None else "" for cell in rows[0]]

    def find_col(*candidates: str) -> int | None:
        for c in candidates:
            if c in header_row:
                return header_row.index(c)
        return None

    code_idx = find_col("service_code", "codigo_servicio", "codigo", "code")
    desc_idx = find_col("description", "descripcion", "concepto")
    unit_idx = find_col("unit", "unidad")
    rate_idx = find_col("allowed_rate", "rate", "tarifa", "tarifa_acordada", "precio")
    curr_idx = find_col("currency", "moneda")

    if code_idx is None or unit_idx is None or rate_idx is None:
        msg = "Encabezados requeridos no encontrados (service_code, unit, allowed_rate)"
        raise HTTPException(400, msg)

    tariffs: list[dict] = []
    for row_idx, row in enumerate(rows[1:], start=2):
        if not any(row):
            continue

        raw_code = row[code_idx]
        raw_unit = row[unit_idx]
        raw_rate = row[rate_idx]
        if isinstance(raw_rate, str) and raw_rate.startswith("="):
            raise HTTPException(422, "La tarifa contiene una fórmula; usa un importe explícito.")
        raw_desc = row[desc_idx] if desc_idx is not None and desc_idx < len(row) else ""
        raw_curr = row[curr_idx] if curr_idx is not None and curr_idx < len(row) else "USD"

        if raw_code is None or raw_rate is None or raw_unit is None:
            continue

        code_str = str(raw_code).strip().upper()
        unit_str = str(raw_unit).strip().upper()
        curr_str = str(raw_curr).strip().upper() if raw_curr else "USD"

        if curr_str != "USD":
            continue

        if unit_str not in ("HOUR", "UNIT"):
            if "HORA" in unit_str or "HR" in unit_str:
                unit_str = "HOUR"
            elif "UNIDAD" in unit_str or "PZA" in unit_str or "PIEZA" in unit_str:
                unit_str = "UNIT"
            else:
                raise HTTPException(422, "Unidad de tarifa desconocida; requiere corrección.")

        try:
            rate_val = Decimal(str(raw_rate).strip()).quantize(
                Decimal("0.01"), rounding=ROUND_HALF_UP
            )
            if rate_val < 0:
                continue
        except (InvalidOperation, ValueError):
            continue

        tariffs.append(
            {
                "service_code": code_str,
                "description": str(raw_desc).strip() or code_str,
                "unit": unit_str,
                "allowed_rate": f"{rate_val:.2f}",
                "location": f"hoja '{sheet.title}', fila {row_idx}",
                "text": f"{code_str}: USD {rate_val:.2f} por {unit_str}",
            }
        )

    return tariffs


LINE_ITEM_REGEX = re.compile(
    r"([A-Za-z0-9áéíóúÁÉÍÓÚñÑ_.\- ]+?):\s*([0-9.]+)\s*"
    r"(HOUR|UNIT|HORAS?|UNIDADES?)\s*[×x*]\s*(?:USD|\$)?\s*([0-9.]+)\s*=\s*(?:USD|\$)?\s*([0-9.]+)",
    re.IGNORECASE,
)


def parse_line_items_from_text(text: str, doc_id: str) -> list[ExtractedItem]:
    """Extrae líneas de facturación a partir del texto del documento de cobro."""
    items: list[ExtractedItem] = []
    lines = [line.strip() for line in text.splitlines() if line.strip()]

    item_counter = 1
    for line in lines:
        for match in LINE_ITEM_REGEX.finditer(line):
            desc = match.group(1).strip()
            qty_str = match.group(2).strip()
            unit_raw = match.group(3).strip().upper()
            unit = "HOUR" if "H" in unit_raw else "UNIT"
            price_str = match.group(4).strip()
            total_str = match.group(5).strip()

            service_code = None
            damage_code = None
            upper_desc = desc.upper()
            if "PINTURA" in upper_desc or "PAINT" in upper_desc:
                service_code = "PAINT"
                damage_code = "FRONT_PAINT"
            elif "PARACHOQUES" in upper_desc or "BUMPER" in upper_desc:
                service_code = "BUMPER"
                damage_code = "FRONT_BUMPER"
            elif "ALINEACI" in upper_desc or "ALIGN" in upper_desc:
                service_code = "ALIGNMENT"
                damage_code = "FRONT_ALIGNMENT"
            elif "SUSPENSI" in upper_desc or "SUSPENSION" in upper_desc:
                service_code = "SUSPENSION"
                damage_code = "SUSPENSION_DAMAGE"

            item_id = f"item_{item_counter}"
            item_counter += 1

            items.append(
                ExtractedItem(
                    id=item_id,
                    description=desc,
                    service_code=service_code,
                    damage_code=damage_code,
                    unit=unit,
                    quantity=Decimal(qty_str),
                    unit_price=Decimal(price_str).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP),
                    line_total=Decimal(total_str).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP),
                    evidence_id=f"ev_line_{item_id}",
                    review_reason=None if service_code else "Código de servicio no identificado",
                )
            )

    return items


def parse_tariffs_from_text(text: str) -> list[dict]:
    """Extrae tarifas desde texto libre o fixtures sintéticos (.txt)."""
    tariffs: list[dict] = []
    pattern = re.compile(
        r"([A-Za-z0-9_-]+):\s*(?:USD|\$)?\s*([0-9.]+)\s*por\s*([A-Za-z]+)",
        re.IGNORECASE,
    )
    for line_idx, line in enumerate(text.splitlines(), start=1):
        line = line.strip()
        if not line:
            continue
        m = pattern.search(line)
        if m:
            code = m.group(1).strip().upper()
            rate_str = m.group(2).strip()
            unit_raw = m.group(3).strip().upper()
            unit = "HOUR" if "H" in unit_raw else "UNIT"
            try:
                rate_val = Decimal(rate_str).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
                tariffs.append(
                    {
                        "service_code": code,
                        "description": f"Tarifa convenida {code}",
                        "unit": unit,
                        "allowed_rate": f"{rate_val:.2f}",
                        "location": f"línea {line_idx}",
                        "text": line,
                    }
                )
            except (InvalidOperation, ValueError):
                continue
    return tariffs


def extract_claim_snapshot(
    claim_id: str,
    active_docs: list[StoredDocument],
    doc_paths: dict[str, Path],
) -> ExtractionSnapshot:
    """Extrae evidencia estructurada y candidatos de todos los documentos activos del claim."""
    from datetime import UTC, datetime

    now = datetime.now(UTC).isoformat()
    evidence_list: list[Evidence] = []
    items: list[ExtractedItem] = []
    tariffs: list[ExtractedTariff] = []
    reported_damages: list[str] = []
    inspected_damages: list[str] = []
    review_notes: list[str] = []

    subtotal: Decimal | None = None
    total: Decimal | None = None
    billing_kind = "INVOICE"

    for doc in active_docs:
        path = doc_paths.get(doc.id)
        if not path or not path.exists():
            review_notes.append(f"Archivo no encontrado en almacenamiento para documento {doc.id}")
            continue

        if doc.kind == DocumentKind.TARIFF:
            is_xlsx = path.suffix.lower() == ".xlsx" or (
                doc.mime.endswith("sheet") and not path.suffix.lower().endswith(".txt")
            )
            if is_xlsx:
                try:
                    extracted_tariffs = extract_xlsx_tariffs(path)
                    for t in extracted_tariffs:
                        ev_id = f"ev_rate_{t['service_code']}"
                        evidence_list.append(
                            Evidence(
                                id=ev_id,
                                document_id=doc.id,
                                location=t["location"],
                                text=t["text"],
                            )
                        )
                        tariffs.append(
                            ExtractedTariff(
                                service_code=t["service_code"],
                                unit=t["unit"],
                                allowed_rate=Decimal(t["allowed_rate"]),
                                evidence_id=ev_id,
                            )
                        )
                except Exception as e:
                    review_notes.append(f"No fue posible leer tarifario Excel {doc.filename}: {e}")
            elif path.suffix.lower() in (".txt", ".text") or doc.mime.startswith("text/"):
                extracted_tariffs = parse_tariffs_from_text(path.read_text(errors="replace"))
                for t in extracted_tariffs:
                    ev_id = f"ev_rate_{t['service_code']}"
                    evidence_list.append(
                        Evidence(
                            id=ev_id,
                            document_id=doc.id,
                            location=t["location"],
                            text=t["text"],
                        )
                    )
                    tariffs.append(
                        ExtractedTariff(
                            service_code=t["service_code"],
                            unit=t["unit"],
                            allowed_rate=Decimal(t["allowed_rate"]),
                            evidence_id=ev_id,
                        )
                    )
            elif path.suffix.lower() == ".pdf":
                pages = extract_pdf_pages(path)
                for p_num, p_text in pages:
                    ev_id = f"ev_pdf_tariff_p{p_num}"
                    evidence_list.append(
                        Evidence(
                            id=ev_id,
                            document_id=doc.id,
                            location=f"página {p_num}",
                            text=p_text[:500],
                        )
                    )

        elif doc.kind in (DocumentKind.INCIDENT, DocumentKind.INSPECTION, DocumentKind.BILLING):
            pages = (
                extract_pdf_pages(path)
                if path.suffix.lower() == ".pdf"
                else [(1, path.read_text(errors="replace"))]
            )

            for p_num, p_text in pages:
                if "[DOCUMENTO ESCANEADO" in p_text:
                    note = f"Doc {doc.kind} ({doc.filename}) sin texto legible (requiere OCR)"
                    review_notes.append(note)
                    continue

                ev_id = f"ev_{doc.kind.value.lower()}_p{p_num}"
                evidence_list.append(
                    Evidence(
                        id=ev_id,
                        document_id=doc.id,
                        location=f"página {p_num}",
                        text=p_text[:1000],
                    )
                )

                upper_text = p_text.upper()
                if doc.kind == DocumentKind.INCIDENT:
                    if "PARACHOQUES" in upper_text or "BUMPER" in upper_text:
                        reported_damages.append("FRONT_BUMPER")
                    if "PINTURA" in upper_text or "PAINT" in upper_text:
                        reported_damages.append("FRONT_PAINT")
                    if "ALINEACI" in upper_text or "ALIGN" in upper_text:
                        reported_damages.append("FRONT_ALIGNMENT")
                    if "SUSPENSI" in upper_text:
                        reported_damages.append("SUSPENSION_DAMAGE")

                elif doc.kind == DocumentKind.INSPECTION:
                    if "PARACHOQUES" in upper_text or "BUMPER" in upper_text:
                        inspected_damages.append("FRONT_BUMPER")
                    if "PINTURA" in upper_text or "PAINT" in upper_text:
                        inspected_damages.append("FRONT_PAINT")
                    if "ALINEACI" in upper_text or "ALIGN" in upper_text:
                        inspected_damages.append("FRONT_ALIGNMENT")

                elif doc.kind == DocumentKind.BILLING:
                    if "COTIZACI" in upper_text or "QUOTE" in upper_text:
                        billing_kind = "QUOTE"
                    parsed_items = parse_line_items_from_text(p_text, doc.id)
                    for item in parsed_items:
                        items.append(item)
                        item_text = next(
                            match.group(0)
                            for match in LINE_ITEM_REGEX.finditer(p_text)
                            if match.group(1).strip() == item.description
                        )
                        evidence_list.append(
                            Evidence(
                                id=item.evidence_id,
                                document_id=doc.id,
                                location=f"página {p_num}",
                                text=item_text,
                            )
                        )

    # Calcular subtotal y total de líneas extraídas
    if items:
        computed_subtotal = sum((i.line_total for i in items if i.line_total), Decimal("0.00"))
        subtotal = computed_subtotal.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
        total = subtotal

    return ExtractionSnapshot(
        claim_id=claim_id,
        created_at=now,
        confirmed=False,
        billing_kind=billing_kind,
        reported_damage_codes=list(dict.fromkeys(reported_damages)),
        inspected_damage_codes=list(dict.fromkeys(inspected_damages)),
        items=items,
        tariffs=tariffs,
        subtotal=subtotal,
        taxes=Decimal("0.00"),
        total=total,
        evidence=evidence_list,
        review_notes=review_notes,
    )
